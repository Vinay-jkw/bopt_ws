#!/usr/bin/env python3

import os
import math
import csv
from datetime import datetime
from std_msgs.msg import String
import rclpy
from rclpy.node import Node

from geometry_msgs.msg import PoseStamped, PoseWithCovarianceStamped, Twist
from nav_msgs.msg import Odometry
from sensor_msgs.msg import Imu, JointState

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


# ============================================================
# Utility functions
# ============================================================

def quaternion_to_yaw(q):
    """
    Convert quaternion to yaw in radians.
    """
    siny_cosp = 2.0 * (q.w * q.z + q.x * q.y)
    cosy_cosp = 1.0 - 2.0 * (q.y * q.y + q.z * q.z)

    return math.atan2(siny_cosp, cosy_cosp)


def wrap_to_pi(angle):
    """
    Wrap angle to [-pi, pi].
    """
    while angle > math.pi:
        angle -= 2.0 * math.pi

    while angle < -math.pi:
        angle += 2.0 * math.pi

    return angle


def distance(x1, y1, x2, y2):
    return math.sqrt(
        (x1 - x2) ** 2 +
        (y1 - y2) ** 2
    )


# ============================================================
# Localization Benchmark Node
# ============================================================

class LocalizationBenchmark(Node):

    def __init__(self):
        super().__init__('localization_benchmark')

        # ----------------------------------------------------
        # Parameters
        # ----------------------------------------------------

        self.declare_parameter(
            'output_directory',
            os.path.expanduser('~/bopt_ws/localization_results')
        )

        self.declare_parameter(
            'record_rate',
            20.0
        )

        output_directory = self.get_parameter(
            'output_directory'
        ).value

        self.record_rate = float(
            self.get_parameter('record_rate').value
        )

        os.makedirs(
            output_directory,
            exist_ok=True
        )

        timestamp = datetime.now().strftime(
            '%Y%m%d_%H%M%S'
        )

        self.output_file = os.path.join(
            output_directory,
            f'localization_test_{timestamp}.xlsx'
        )

        # ----------------------------------------------------
        # Latest measurements
        # ----------------------------------------------------

        self.gt = None
        self.odom = None
        self.ekf = None
        self.amcl = None
        self.imu_yaw = None
        self.imu_yaw_rate = 0.0

        # Command state
        self.cmd_vehicle_velocity = 0.0
        self.cmd_steering_command = 0.0

        # Joint state
        self.drive_wheel_position = 0.0
        self.drive_wheel_velocity = 0.0
        self.steering_position = 0.0
        self.steering_velocity = 0.0

        # ----------------------------------------------------
        # Start reference
        # ----------------------------------------------------

        self.start_gt = None
        self.start_odom = None
        self.start_ekf = None
        self.start_amcl = None
        self.start_imu_yaw = None
        self.phase = 'IDLE'
        self.started = False

        # ----------------------------------------------------
        # Data storage
        # ----------------------------------------------------

        self.raw_data = []

        # ----------------------------------------------------
        # Subscribers
        # ----------------------------------------------------

        self.create_subscription(
            PoseStamped,
            '/ground_truth/pose',
            self.ground_truth_callback,
            20
        )

        self.create_subscription(
            Odometry,
            '/odom',
            self.odom_callback,
            20
        )

        self.create_subscription(
            Odometry,
            '/odometry/filtered',
            self.ekf_callback,
            20
        )

        self.create_subscription(
            PoseWithCovarianceStamped,
            '/amcl_pose',
            self.amcl_callback,
            20
        )

        self.create_subscription(
            Imu,
            '/imu/corrected',
            self.imu_callback,
            20
        )

        self.create_subscription(
            Twist,
            '/cmd_vel',
            self.cmd_vel_callback,
            20
        )

        self.create_subscription(
            JointState,
            '/joint_states',
            self.joint_state_callback,
            20
        )

        self.create_subscription(
            String,
            '/localization_test/phase',
            self.phase_callback,
            20
        )

        # ----------------------------------------------------
        # Recording timer
        # ----------------------------------------------------

        timer_period = 1.0 / self.record_rate

        self.timer = self.create_timer(
            timer_period,
            self.record_callback
        )

        self.get_logger().info(
            '=============================================='
        )

        self.get_logger().info(
            'BOPT LOCALIZATION BENCHMARK'
        )

        self.get_logger().info(
            '=============================================='
        )

        self.get_logger().info(
            f'Record rate: {self.record_rate:.1f} Hz'
        )

        self.get_logger().info(
            f'Output: {self.output_file}'
        )

        self.get_logger().info(
            'Waiting for all localization sources...'
        )

    # ========================================================
    # Callbacks
    # ========================================================

    def ground_truth_callback(self, msg):

        self.gt = {
            'x': msg.pose.position.x,
            'y': msg.pose.position.y,
            'z': msg.pose.position.z,
            'yaw': quaternion_to_yaw(
                msg.pose.orientation
            )
        }

    def odom_callback(self, msg):

        self.odom = {
            'x': msg.pose.pose.position.x,
            'y': msg.pose.pose.position.y,
            'z': msg.pose.pose.position.z,
            'yaw': quaternion_to_yaw(
                msg.pose.pose.orientation
            )
        }

    def ekf_callback(self, msg):

        self.ekf = {
            'x': msg.pose.pose.position.x,
            'y': msg.pose.pose.position.y,
            'z': msg.pose.pose.position.z,
            'yaw': quaternion_to_yaw(
                msg.pose.pose.orientation
            )
        }

    def amcl_callback(self, msg):

        self.amcl = {
            'x': msg.pose.pose.position.x,
            'y': msg.pose.pose.position.y,
            'z': msg.pose.pose.position.z,
            'yaw': quaternion_to_yaw(
                msg.pose.pose.orientation
            )
        }

    def imu_callback(self, msg):

        self.imu_yaw = quaternion_to_yaw(
            msg.orientation
        )
        self.imu_yaw_rate = msg.angular_velocity.z

    def cmd_vel_callback(self, msg):

        self.cmd_vehicle_velocity = msg.linear.x
        self.cmd_steering_command = msg.angular.z

    def joint_state_callback(self, msg):

        for i, name in enumerate(msg.name):
            if name == 'drive_wheel_joint':
                if i < len(msg.position):
                    self.drive_wheel_position = msg.position[i]
                if i < len(msg.velocity):
                    self.drive_wheel_velocity = msg.velocity[i]
            elif name in ['drive_wheel_Ass_joint', 'steering_joint']:
                if i < len(msg.position):
                    self.steering_position = msg.position[i]
                if i < len(msg.velocity):
                    self.steering_velocity = msg.velocity[i]

    def phase_callback(self, msg):

        self.phase = msg.data

    # ========================================================
    # Check whether all sources are available
    # ========================================================

    def all_sources_ready(self):

        return (
            self.gt is not None and
            self.odom is not None and
            self.ekf is not None and
            self.amcl is not None and
            self.imu_yaw is not None
        )

    # ========================================================
    # Establish START reference
    # ========================================================

    def establish_start(self):

        self.start_gt = self.gt.copy()
        self.start_odom = self.odom.copy()
        self.start_ekf = self.ekf.copy()
        self.start_amcl = self.amcl.copy()
        self.start_imu_yaw = self.imu_yaw

        self.started = True

        self.get_logger().info(
            '=============================================='
        )

        self.get_logger().info(
            'START REFERENCE CAPTURED'
        )

        self.get_logger().info(
            f"GT    : x={self.start_gt['x']:.4f}, "
            f"y={self.start_gt['y']:.4f}, "
            f"yaw={math.degrees(self.start_gt['yaw']):.2f} deg"
        )

        self.get_logger().info(
            f"Odom  : x={self.start_odom['x']:.4f}, "
            f"y={self.start_odom['y']:.4f}, "
            f"yaw={math.degrees(self.start_odom['yaw']):.2f} deg"
        )

        self.get_logger().info(
            f"EKF   : x={self.start_ekf['x']:.4f}, "
            f"y={self.start_ekf['y']:.4f}, "
            f"yaw={math.degrees(self.start_ekf['yaw']):.2f} deg"
        )

        self.get_logger().info(
            f"AMCL  : x={self.start_amcl['x']:.4f}, "
            f"y={self.start_amcl['y']:.4f}, "
            f"yaw={math.degrees(self.start_amcl['yaw']):.2f} deg"
        )

        self.get_logger().info(
            f"IMU   : yaw={math.degrees(self.start_imu_yaw):.2f} deg"
        )

        self.get_logger().info(
            '=============================================='
        )

    # ========================================================
    # Calculate errors
    # ========================================================

    def calculate_error(
        self,
        estimate,
        ground_truth
    ):

        dx = estimate['x'] - ground_truth['x']
        dy = estimate['y'] - ground_truth['y']

        position_error = math.sqrt(
            dx * dx +
            dy * dy
        )

        yaw_error = wrap_to_pi(
            estimate['yaw'] -
            ground_truth['yaw']
        )

        return (
            dx,
            dy,
            position_error,
            yaw_error
        )

    # ========================================================
    # Start-aligned displacement
    # ========================================================

    def aligned_error(
        self,
        estimate,
        start_estimate,
        ground_truth,
        start_ground_truth
    ):

        gt_dx = (
            ground_truth['x'] -
            start_ground_truth['x']
        )

        gt_dy = (
            ground_truth['y'] -
            start_ground_truth['y']
        )

        est_dx = (
            estimate['x'] -
            start_estimate['x']
        )

        est_dy = (
            estimate['y'] -
            start_estimate['y']
        )

        dx = est_dx - gt_dx
        dy = est_dy - gt_dy

        error = math.sqrt(
            dx * dx +
            dy * dy
        )

        return (
            gt_dx,
            gt_dy,
            est_dx,
            est_dy,
            dx,
            dy,
            error
        )

    # ========================================================
    # Recording
    # ========================================================

    def record_callback(self):

        if not self.all_sources_ready():

            return

        if not self.started:

            self.establish_start()

            return

        now = self.get_clock().now()

        timestamp_ns = now.nanoseconds

        timestamp_sec = timestamp_ns / 1e9

        # ----------------------------------------------------
        # Position / heading errors
        # ----------------------------------------------------

        (
            odom_dx,
            odom_dy,
            odom_pos_error,
            odom_yaw_error
        ) = self.calculate_error(
            self.odom,
            self.gt
        )

        (
            ekf_dx,
            ekf_dy,
            ekf_pos_error,
            ekf_yaw_error
        ) = self.calculate_error(
            self.ekf,
            self.gt
        )

        (
            amcl_dx,
            amcl_dy,
            amcl_pos_error,
            amcl_yaw_error
        ) = self.calculate_error(
            self.amcl,
            self.gt
        )

        imu_yaw_error = wrap_to_pi(
            self.imu_yaw -
            self.gt['yaw']
        )

        # ----------------------------------------------------
        # Start-aligned errors
        # ----------------------------------------------------

        (
            gt_dx_odom,
            gt_dy_odom,
            est_dx_odom,
            est_dy_odom,
            aligned_odom_dx,
            aligned_odom_dy,
            aligned_odom_error
        ) = self.aligned_error(
            self.odom,
            self.start_odom,
            self.gt,
            self.start_gt
        )

        (
            gt_dx_ekf,
            gt_dy_ekf,
            est_dx_ekf,
            est_dy_ekf,
            aligned_ekf_dx,
            aligned_ekf_dy,
            aligned_ekf_error
        ) = self.aligned_error(
            self.ekf,
            self.start_ekf,
            self.gt,
            self.start_gt
        )

        (
            gt_dx_amcl,
            gt_dy_amcl,
            est_dx_amcl,
            est_dy_amcl,
            aligned_amcl_dx,
            aligned_amcl_dy,
            aligned_amcl_error
        ) = self.aligned_error(
            self.amcl,
            self.start_amcl,
            self.gt,
            self.start_gt
        )

        # ----------------------------------------------------
        # Store row
        # ----------------------------------------------------

        row = {

            'timestamp_sec': timestamp_sec,

            # COMMAND
            'cmd_vehicle_velocity': self.cmd_vehicle_velocity,
            'cmd_steering_command': self.cmd_steering_command,

            # JOINT STATES
            'drive_wheel_position': self.drive_wheel_position,
            'drive_wheel_velocity': self.drive_wheel_velocity,
            'steering_position': self.steering_position,
            'steering_velocity': self.steering_velocity,

            # Ground truth
            'gt_x': self.gt['x'],
            'gt_y': self.gt['y'],
            'gt_z': self.gt['z'],
            'gt_yaw_rad': self.gt['yaw'],
            'gt_yaw_deg': math.degrees(
                self.gt['yaw']
            ),

            # Odom
            'odom_x': self.odom['x'],
            'odom_y': self.odom['y'],
            'odom_yaw_rad': self.odom['yaw'],
            'odom_yaw_deg': math.degrees(
                self.odom['yaw']
            ),

            'odom_dx': odom_dx,
            'odom_dy': odom_dy,
            'odom_position_error': odom_pos_error,
            'odom_yaw_error_rad': odom_yaw_error,
            'odom_yaw_error_deg': math.degrees(
                odom_yaw_error
            ),

            'odom_aligned_dx': aligned_odom_dx,
            'odom_aligned_dy': aligned_odom_dy,
            'odom_aligned_position_error':
                aligned_odom_error,

            # EKF
            'ekf_x': self.ekf['x'],
            'ekf_y': self.ekf['y'],
            'ekf_yaw_rad': self.ekf['yaw'],
            'ekf_yaw_deg': math.degrees(
                self.ekf['yaw']
            ),

            'ekf_dx': ekf_dx,
            'ekf_dy': ekf_dy,
            'ekf_position_error': ekf_pos_error,
            'ekf_yaw_error_rad': ekf_yaw_error,
            'ekf_yaw_error_deg': math.degrees(
                ekf_yaw_error
            ),

            'ekf_aligned_dx': aligned_ekf_dx,
            'ekf_aligned_dy': aligned_ekf_dy,
            'ekf_aligned_position_error':
                aligned_ekf_error,

            # AMCL
            'amcl_x': self.amcl['x'],
            'amcl_y': self.amcl['y'],
            'amcl_yaw_rad': self.amcl['yaw'],
            'amcl_yaw_deg': math.degrees(
                self.amcl['yaw']
            ),

            'amcl_dx': amcl_dx,
            'amcl_dy': amcl_dy,
            'amcl_position_error': amcl_pos_error,
            'amcl_yaw_error_rad': amcl_yaw_error,
            'amcl_yaw_error_deg': math.degrees(
                amcl_yaw_error
            ),

            'amcl_aligned_dx': aligned_amcl_dx,
            'amcl_aligned_dy': aligned_amcl_dy,
            'amcl_aligned_position_error':
                aligned_amcl_error,

            # Phase
            'phase': self.phase,

            # IMU
            'imu_yaw_rad': self.imu_yaw,
            'imu_yaw_deg': math.degrees(
                self.imu_yaw
            ),
            'imu_yaw_rate_rad_s': self.imu_yaw_rate,
            'imu_yaw_rate_deg_s': math.degrees(
                self.imu_yaw_rate
            ),

            'imu_yaw_error_rad': imu_yaw_error,
            'imu_yaw_error_deg': math.degrees(
                imu_yaw_error
            ),
        }

        self.raw_data.append(row)

        # ----------------------------------------------------
        # Terminal status
        # ----------------------------------------------------

        if len(self.raw_data) % int(
            self.record_rate * 2
        ) == 0:

            self.get_logger().info(
                f"[{self.phase:^14}] GT=({self.gt['x']:6.3f}, {self.gt['y']:6.3f}) | "
                f"OdomErr={aligned_odom_error:6.3f}m | "
                f"EKFErr={aligned_ekf_error:6.3f}m | "
                f"AMCLErr={aligned_amcl_error:6.3f}m | "
                f"IMU yaw={math.degrees(imu_yaw_error):6.2f}deg"
            )

    # ========================================================
    # Generate Excel report
    # ========================================================

    def save_results(self):

        if not self.raw_data:

            self.get_logger().warn(
                'No data recorded.'
            )

            return

        self.get_logger().info(
            'Generating Excel report...'
        )

        df = pd.DataFrame(
            self.raw_data
        )

        # ----------------------------------------------------
        # Summary statistics
        # ----------------------------------------------------

        summary_rows = []

        sources = [
            (
                'Odom',
                'odom_position_error',
                'odom_yaw_error_deg',
                'odom_aligned_position_error'
            ),
            (
                'EKF',
                'ekf_position_error',
                'ekf_yaw_error_deg',
                'ekf_aligned_position_error'
            ),
            (
                'AMCL',
                'amcl_position_error',
                'amcl_yaw_error_deg',
                'amcl_aligned_position_error'
            ),
        ]

        for name, pos_col, yaw_col, aligned_col in sources:

            position_values = df[pos_col].dropna()

            yaw_values = df[yaw_col].dropna()

            aligned_values = df[
                aligned_col
            ].dropna()

            if len(position_values) == 0:

                continue

            summary_rows.append({

                'Source': name,

                'Samples': len(position_values),

                'Mean Position Error (m)':
                    position_values.mean(),

                'RMSE Position Error (m)':
                    math.sqrt(
                        (
                            position_values ** 2
                        ).mean()
                    ),

                'Maximum Position Error (m)':
                    position_values.max(),

                'Final Position Error (m)':
                    position_values.iloc[-1],

                'Mean Aligned Position Error (m)':
                    aligned_values.mean(),

                'RMSE Aligned Position Error (m)':
                    math.sqrt(
                        (
                            aligned_values ** 2
                        ).mean()
                    ),

                'Maximum Aligned Position Error (m)':
                    aligned_values.max(),

                'Final Aligned Position Error (m)':
                    aligned_values.iloc[-1],

                'Mean Heading Error (deg)':
                    yaw_values.abs().mean(),

                'RMSE Heading Error (deg)':
                    math.sqrt(
                        (
                            yaw_values ** 2
                        ).mean()
                    ),

                'Maximum Heading Error (deg)':
                    yaw_values.abs().max(),

                'Final Heading Error (deg)':
                    abs(yaw_values.iloc[-1])
            })

        # ----------------------------------------------------
        # IMU summary
        # ----------------------------------------------------

        imu_values = df[
            'imu_yaw_error_deg'
        ].dropna()

        if len(imu_values) > 0:

            summary_rows.append({

                'Source': 'IMU',

                'Samples': len(imu_values),

                'Mean Position Error (m)':
                    None,

                'RMSE Position Error (m)':
                    None,

                'Maximum Position Error (m)':
                    None,

                'Final Position Error (m)':
                    None,

                'Mean Aligned Position Error (m)':
                    None,

                'RMSE Aligned Position Error (m)':
                    None,

                'Maximum Aligned Position Error (m)':
                    None,

                'Final Aligned Position Error (m)':
                    None,

                'Mean Heading Error (deg)':
                    imu_values.abs().mean(),

                'RMSE Heading Error (deg)':
                    math.sqrt(
                        (
                            imu_values ** 2
                        ).mean()
                    ),

                'Maximum Heading Error (deg)':
                    imu_values.abs().max(),

                'Final Heading Error (deg)':
                    abs(imu_values.iloc[-1])
            })

        summary_df = pd.DataFrame(
            summary_rows
        )

        # ----------------------------------------------------
        # Checkpoint: final state
        # ----------------------------------------------------

        final = df.iloc[-1]

        checkpoint_df = pd.DataFrame([{

            'Checkpoint': 'FINAL',

            'Timestamp (s)':
                final['timestamp_sec'],

            'Cmd Vehicle Velocity (m/s)':
                final['cmd_vehicle_velocity'],

            'Cmd Steering Command (rad/s)':
                final['cmd_steering_command'],

            'Drive Wheel Position (rad)':
                final['drive_wheel_position'],

            'Drive Wheel Velocity (rad/s)':
                final['drive_wheel_velocity'],

            'Steering Position (rad)':
                final['steering_position'],

            'Steering Velocity (rad/s)':
                final['steering_velocity'],

            'GT X (m)':
                final['gt_x'],

            'GT Y (m)':
                final['gt_y'],

            'GT Yaw (deg)':
                final['gt_yaw_deg'],

            'Odom X (m)':
                final['odom_x'],

            'Odom Y (m)':
                final['odom_y'],

            'Odom Yaw (deg)':
                final['odom_yaw_deg'],

            'Odom Position Error (m)':
                final['odom_position_error'],

            'Odom Heading Error (deg)':
                final['odom_yaw_error_deg'],

            'EKF X (m)':
                final['ekf_x'],

            'EKF Y (m)':
                final['ekf_y'],

            'EKF Yaw (deg)':
                final['ekf_yaw_deg'],

            'EKF Position Error (m)':
                final['ekf_position_error'],

            'EKF Heading Error (deg)':
                final['ekf_yaw_error_deg'],

            'AMCL X (m)':
                final['amcl_x'],

            'AMCL Y (m)':
                final['amcl_y'],

            'AMCL Yaw (deg)':
                final['amcl_yaw_deg'],

            'AMCL Position Error (m)':
                final['amcl_position_error'],

            'AMCL Heading Error (deg)':
                final['amcl_yaw_error_deg'],

            'IMU Yaw Rate (deg/s)':
                final['imu_yaw_rate_deg_s'],

            'IMU Heading Error (deg)':
                final['imu_yaw_error_deg']
        }])

        # ----------------------------------------------------
        # Heading error sheet
        # ----------------------------------------------------

        heading_df = df[[
            'timestamp_sec',
            'gt_yaw_deg',
            'odom_yaw_deg',
            'odom_yaw_error_deg',
            'ekf_yaw_deg',
            'ekf_yaw_error_deg',
            'amcl_yaw_deg',
            'amcl_yaw_error_deg',
            'imu_yaw_deg',
            'imu_yaw_error_deg'
        ]].copy()

        # ----------------------------------------------------
        # Command & Joint States sheet
        # ----------------------------------------------------

        cmd_joints_df = df[[
            'timestamp_sec',
            'cmd_vehicle_velocity',
            'cmd_steering_command',
            'drive_wheel_position',
            'drive_wheel_velocity',
            'steering_position',
            'steering_velocity',
            'imu_yaw_rate_deg_s',
            'phase'
        ]].copy()

        # ----------------------------------------------------
        # Position error sheet
        # ----------------------------------------------------

        position_df = df[[
            'timestamp_sec',

            'gt_x',
            'gt_y',

            'odom_dx',
            'odom_dy',
            'odom_position_error',

            'ekf_dx',
            'ekf_dy',
            'ekf_position_error',

            'amcl_dx',
            'amcl_dy',
            'amcl_position_error',

            'odom_aligned_dx',
            'odom_aligned_dy',
            'odom_aligned_position_error',

            'ekf_aligned_dx',
            'ekf_aligned_dy',
            'ekf_aligned_position_error',

            'amcl_aligned_dx',
            'amcl_aligned_dy',
            'amcl_aligned_position_error'
        ]].copy()

        # ----------------------------------------------------
        # Write Excel
        # ----------------------------------------------------

        with pd.ExcelWriter(
            self.output_file,
            engine='openpyxl'
        ) as writer:

            summary_df.to_excel(
                writer,
                sheet_name='Summary',
                index=False
            )

            df.to_excel(
                writer,
                sheet_name='Raw_Data',
                index=False
            )

            cmd_joints_df.to_excel(
                writer,
                sheet_name='Command_Joints',
                index=False
            )

            position_df.to_excel(
                writer,
                sheet_name='Position_Error',
                index=False
            )

            heading_df.to_excel(
                writer,
                sheet_name='Heading_Error',
                index=False
            )

            checkpoint_df.to_excel(
                writer,
                sheet_name='Checkpoints',
                index=False
            )

        # ----------------------------------------------------
        # Format workbook
        # ----------------------------------------------------

        workbook = load_workbook(
            self.output_file
        )

        for worksheet in workbook.worksheets:

            worksheet.freeze_panes = 'A2'

            worksheet.auto_filter.ref = (
                worksheet.dimensions
            )

            for cell in worksheet[1]:

                cell.font = Font(
                    bold=True
                )

            for column_cells in worksheet.columns:

                max_length = 0

                column_letter = get_column_letter(
                    column_cells[0].column
                )

                for cell in column_cells:

                    try:
                        value_length = len(
                            str(cell.value)
                        )

                        max_length = max(
                            max_length,
                            value_length
                        )

                    except Exception:
                        pass

                worksheet.column_dimensions[
                    column_letter
                ].width = min(
                    max_length + 2,
                    30
                )

        workbook.save(
            self.output_file
        )

        try:
            self.get_logger().info(
                '=============================================='
            )
            self.get_logger().info(
                'LOCALIZATION TEST COMPLETE'
            )
            self.get_logger().info(
                f'Samples recorded: {len(df)}'
            )
            self.get_logger().info(
                f'Excel file: {self.output_file}'
            )
            self.get_logger().info(
                '=============================================='
            )
        except Exception:
            print('==============================================')
            print('LOCALIZATION TEST COMPLETE')
            print(f'Samples recorded: {len(df)}')
            print(f'Excel file: {self.output_file}')
            print('==============================================')


# ============================================================
# Main
# ============================================================

def main(args=None):

    rclpy.init(args=args)

    node = LocalizationBenchmark()

    try:

        rclpy.spin(node)

    except KeyboardInterrupt:

        try:
            node.get_logger().info(
                'Stopping benchmark...'
            )
        except Exception:
            print('Stopping benchmark...')

    finally:

        try:
            node.save_results()
        except Exception as e:
            print(f'Error saving results: {e}')

        try:
            node.destroy_node()
        except Exception:
            pass

        if rclpy.ok():
            try:
                rclpy.shutdown()
            except Exception:
                pass


if __name__ == '__main__':
    main()
