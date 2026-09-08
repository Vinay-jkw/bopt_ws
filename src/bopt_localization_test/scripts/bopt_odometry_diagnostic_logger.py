#!/usr/bin/env python3
"""
BOPT Odometry Diagnostic Logger V2
==================================

Purpose:
    Diagnose encoder-based odometry without changing the robot/controller.

Key corrections from V1:
    1. Wheel command is treated as rad/s, matching the BOPT controller.
    2. Actual wheel velocity is compared directly with wheel command.
    3. ROS message timestamps are kept separate from wall-clock time.
    4. Ground-truth topic type can be selected:
         --gt-type pose_stamped
         --gt-type odometry
    5. Timestamp age is only calculated when both timestamps use the same
       ROS clock domain.
    6. Raw JointState data is preserved.
    7. Encoder-only diagnostics are calculated independently of GT.
    8. GT-based calibration is calculated only when valid GT data exists.

This script DOES NOT command the robot.
The operator performs each test motion.

Default phases:
    stationary
    forward_straight
    reverse_straight
    forward_left
    forward_right
    reverse_left
    reverse_right
    forward_reverse_transition
    reverse_forward_transition

Excel sheets:
    Raw_Data
    Phase_Summary
    Encoder_Diagnostics
    Curvature_Diagnostics
    Calibration
    Configuration

Run:
    source ~/bopt_ws/install/setup.bash
    python3 bopt_odometry_diagnostic_logger_v2.py

Before running, verify the GT topic:
    ros2 topic list | grep ground_truth
    ros2 topic info /ground_truth/pose

If GT is nav_msgs/Odometry:
    --gt-type odometry

If GT is geometry_msgs/PoseStamped:
    --gt-type pose_stamped
"""

import argparse
import math
import time
from dataclasses import dataclass
from pathlib import Path

import rclpy
from rclpy.node import Node
from rclpy.qos import (
    QoSProfile,
    ReliabilityPolicy,
    DurabilityPolicy,
    HistoryPolicy,
)

from sensor_msgs.msg import JointState, Imu
from nav_msgs.msg import Odometry
from geometry_msgs.msg import PoseStamped, PoseWithCovarianceStamped
from std_msgs.msg import Float64MultiArray

from openpyxl import Workbook
from openpyxl.utils import get_column_letter


def NaN():
    return float("nan")


def finite(x):
    try:
        return math.isfinite(float(x))
    except Exception:
        return False


def stamp_sec(stamp):
    return float(stamp.sec) + float(stamp.nanosec) * 1e-9


def yaw_from_quaternion(q):
    return math.atan2(
        2.0 * (q.w * q.z + q.x * q.y),
        1.0 - 2.0 * (q.y * q.y + q.z * q.z),
    )


def wrap_angle(a):
    return (a + math.pi) % (2.0 * math.pi) - math.pi


def mean(values):
    values = [float(v) for v in values if finite(v)]
    return sum(values) / len(values) if values else NaN()


def max_abs(values):
    values = [abs(float(v)) for v in values if finite(v)]
    return max(values) if values else NaN()


def rms(values):
    values = [float(v) for v in values if finite(v)]
    return math.sqrt(sum(v * v for v in values) / len(values)) if values else NaN()


@dataclass
class Row:
    wall_time: float
    phase: str

    # GT
    gt_stamp: float
    gt_x: float
    gt_y: float
    gt_yaw: float

    # Odom
    odom_stamp: float
    odom_x: float
    odom_y: float
    odom_yaw: float
    odom_vx: float
    odom_wz: float

    # EKF
    ekf_stamp: float
    ekf_x: float
    ekf_y: float
    ekf_yaw: float
    ekf_vx: float
    ekf_wz: float

    # AMCL
    amcl_stamp: float
    amcl_x: float
    amcl_y: float
    amcl_yaw: float

    # IMU
    imu_stamp: float
    imu_yaw: float
    imu_wz: float

    # JointState
    joint_stamp: float
    drive_pos: float
    drive_vel: float
    steer_pos: float
    steer_vel: float

    # Controller command
    wheel_cmd_radps: float
    steer_cmd_rad: float


class DiagnosticNode(Node):
    def __init__(self, args):
        super().__init__("bopt_odometry_diagnostic_logger_v2")

        self.args = args
        self.phase = "not_started"
        self.rows = []

        self.drive_joint = args.drive_joint
        self.steer_joint = args.steer_joint
        self.R = args.wheel_radius
        self.L = args.wheelbase

        self.latest = {
            "gt": None,
            "odom": None,
            "ekf": None,
            "amcl": None,
            "imu": None,
            "joint": None,
            "wheel_cmd": None,
            "steer_cmd": None,
        }

        self.previous = None

        sensor_qos = QoSProfile(
            history=HistoryPolicy.KEEP_LAST,
            depth=100,
            reliability=ReliabilityPolicy.BEST_EFFORT,
            durability=DurabilityPolicy.VOLATILE,
        )

        reliable_qos = QoSProfile(
            history=HistoryPolicy.KEEP_LAST,
            depth=100,
            reliability=ReliabilityPolicy.RELIABLE,
            durability=DurabilityPolicy.VOLATILE,
        )

        amcl_qos = QoSProfile(
            history=HistoryPolicy.KEEP_LAST,
            depth=20,
            reliability=ReliabilityPolicy.RELIABLE,
            durability=DurabilityPolicy.TRANSIENT_LOCAL,
        )

        self.create_subscription(
            JointState, args.joint_topic, self.joint_cb, sensor_qos
        )
        self.create_subscription(
            Odometry, args.odom_topic, self.odom_cb, reliable_qos
        )
        self.create_subscription(
            Odometry, args.ekf_topic, self.ekf_cb, reliable_qos
        )
        self.create_subscription(
            PoseWithCovarianceStamped, args.amcl_topic, self.amcl_cb, amcl_qos
        )
        self.create_subscription(
            Imu, args.imu_topic, self.imu_cb, sensor_qos
        )
        self.create_subscription(
            Float64MultiArray,
            args.wheel_cmd_topic,
            self.wheel_cmd_cb,
            reliable_qos,
        )
        self.create_subscription(
            Float64MultiArray,
            args.steer_cmd_topic,
            self.steer_cmd_cb,
            reliable_qos,
        )

        if args.gt_type == "odometry":
            self.create_subscription(
                Odometry, args.gt_topic, self.gt_odom_cb, reliable_qos
            )
        else:
            self.create_subscription(
                PoseStamped, args.gt_topic, self.gt_pose_cb, reliable_qos
            )

        self.timer = self.create_timer(1.0 / args.sample_hz, self.sample)

        self.get_logger().info("BOPT Odometry Diagnostic Logger V2 started")
        self.get_logger().info(f"GT type: {args.gt_type}")
        self.get_logger().info(f"GT topic: {args.gt_topic}")
        self.get_logger().info(f"Drive joint: {self.drive_joint}")
        self.get_logger().info(f"Steer joint: {self.steer_joint}")
        self.get_logger().info(f"Nominal wheel radius: {self.R:.6f} m")
        self.get_logger().info(f"Nominal wheelbase: {self.L:.6f} m")

    def joint_cb(self, msg):
        if self.drive_joint not in msg.name or self.steer_joint not in msg.name:
            return

        di = msg.name.index(self.drive_joint)
        si = msg.name.index(self.steer_joint)

        with_default = lambda arr, idx: float(arr[idx]) if idx < len(arr) else NaN()

        self.latest["joint"] = {
            "stamp": stamp_sec(msg.header.stamp),
            "drive_pos": with_default(msg.position, di),
            "drive_vel": with_default(msg.velocity, di),
            "steer_pos": with_default(msg.position, si),
            "steer_vel": with_default(msg.velocity, si),
        }

    def odom_cb(self, msg):
        self.latest["odom"] = {
            "stamp": stamp_sec(msg.header.stamp),
            "x": msg.pose.pose.position.x,
            "y": msg.pose.pose.position.y,
            "yaw": yaw_from_quaternion(msg.pose.pose.orientation),
            "vx": msg.twist.twist.linear.x,
            "wz": msg.twist.twist.angular.z,
        }

    def ekf_cb(self, msg):
        self.latest["ekf"] = {
            "stamp": stamp_sec(msg.header.stamp),
            "x": msg.pose.pose.position.x,
            "y": msg.pose.pose.position.y,
            "yaw": yaw_from_quaternion(msg.pose.pose.orientation),
            "vx": msg.twist.twist.linear.x,
            "wz": msg.twist.twist.angular.z,
        }

    def amcl_cb(self, msg):
        self.latest["amcl"] = {
            "stamp": stamp_sec(msg.header.stamp),
            "x": msg.pose.pose.position.x,
            "y": msg.pose.pose.position.y,
            "yaw": yaw_from_quaternion(msg.pose.pose.orientation),
        }

    def imu_cb(self, msg):
        orientation_valid = (
            abs(msg.orientation.w)
            + abs(msg.orientation.x)
            + abs(msg.orientation.y)
            + abs(msg.orientation.z)
            > 1e-9
        )

        self.latest["imu"] = {
            "stamp": stamp_sec(msg.header.stamp),
            "yaw": yaw_from_quaternion(msg.orientation) if orientation_valid else NaN(),
            "wz": msg.angular_velocity.z,
        }

    def gt_pose_cb(self, msg):
        self.latest["gt"] = {
            "stamp": stamp_sec(msg.header.stamp),
            "x": msg.pose.position.x,
            "y": msg.pose.position.y,
            "yaw": yaw_from_quaternion(msg.pose.orientation),
        }

    def gt_odom_cb(self, msg):
        self.latest["gt"] = {
            "stamp": stamp_sec(msg.header.stamp),
            "x": msg.pose.pose.position.x,
            "y": msg.pose.pose.position.y,
            "yaw": yaw_from_quaternion(msg.pose.pose.orientation),
        }

    def wheel_cmd_cb(self, msg):
        if msg.data:
            self.latest["wheel_cmd"] = {
                "value": float(msg.data[0]),
                # This is receipt time only; command message has no ROS stamp.
                "receipt": time.monotonic(),
            }

    def steer_cmd_cb(self, msg):
        if msg.data:
            self.latest["steer_cmd"] = {
                "value": float(msg.data[0]),
                "receipt": time.monotonic(),
            }

    def sample(self):
        j = self.latest["joint"]
        if j is None:
            return

        gt = self.latest["gt"]
        od = self.latest["odom"]
        ek = self.latest["ekf"]
        am = self.latest["amcl"]
        im = self.latest["imu"]
        wc = self.latest["wheel_cmd"]
        sc = self.latest["steer_cmd"]

        r = Row(
            wall_time=time.monotonic(),
            phase=self.phase,

            gt_stamp=gt["stamp"] if gt else NaN(),
            gt_x=gt["x"] if gt else NaN(),
            gt_y=gt["y"] if gt else NaN(),
            gt_yaw=gt["yaw"] if gt else NaN(),

            odom_stamp=od["stamp"] if od else NaN(),
            odom_x=od["x"] if od else NaN(),
            odom_y=od["y"] if od else NaN(),
            odom_yaw=od["yaw"] if od else NaN(),
            odom_vx=od["vx"] if od else NaN(),
            odom_wz=od["wz"] if od else NaN(),

            ekf_stamp=ek["stamp"] if ek else NaN(),
            ekf_x=ek["x"] if ek else NaN(),
            ekf_y=ek["y"] if ek else NaN(),
            ekf_yaw=ek["yaw"] if ek else NaN(),
            ekf_vx=ek["vx"] if ek else NaN(),
            ekf_wz=ek["wz"] if ek else NaN(),

            amcl_stamp=am["stamp"] if am else NaN(),
            amcl_x=am["x"] if am else NaN(),
            amcl_y=am["y"] if am else NaN(),
            amcl_yaw=am["yaw"] if am else NaN(),

            imu_stamp=im["stamp"] if im else NaN(),
            imu_yaw=im["yaw"] if im else NaN(),
            imu_wz=im["wz"] if im else NaN(),

            joint_stamp=j["stamp"],
            drive_pos=j["drive_pos"],
            drive_vel=j["drive_vel"],
            steer_pos=j["steer_pos"],
            steer_vel=j["steer_vel"],

            wheel_cmd_radps=wc["value"] if wc else NaN(),
            steer_cmd_rad=sc["value"] if sc else NaN(),
        )

        self.rows.append(r)

    def set_phase(self, phase):
        self.phase = phase
        self.get_logger().info(f"PHASE = {phase}")

    def derive(self):
        out = []

        for i, r in enumerate(self.rows):
            prev = self.rows[i - 1] if i > 0 else None

            if prev and finite(r.joint_stamp) and finite(prev.joint_stamp):
                dtj = r.joint_stamp - prev.joint_stamp
            else:
                dtj = NaN()

            valid_dtj = finite(dtj) and dtj > 0

            if valid_dtj:
                dtheta = r.drive_pos - prev.drive_pos
                dsteer = wrap_angle(r.steer_pos - prev.steer_pos)
                distance = self.R * dtheta
                encoder_v = distance / dtj
                steer_rate = dsteer / dtj
                steer_mid = wrap_angle(
                    prev.steer_pos + 0.5 * dsteer
                )
            else:
                dtheta = dsteer = distance = encoder_v = steer_rate = steer_mid = NaN()

            # Correct command comparison:
            # controller publishes wheel angular velocity [rad/s].
            wheel_tracking_error = (
                r.drive_vel - r.wheel_cmd_radps
                if finite(r.drive_vel) and finite(r.wheel_cmd_radps)
                else NaN()
            )

            steering_tracking_error = (
                wrap_angle(r.steer_pos - r.steer_cmd_rad)
                if finite(r.steer_pos) and finite(r.steer_cmd_rad)
                else NaN()
            )

            if finite(encoder_v) and finite(steer_mid):
                model_wz = -encoder_v * math.sin(steer_mid) / self.L
            else:
                model_wz = NaN()

            # GT derivatives
            if prev and finite(r.gt_stamp) and finite(prev.gt_stamp):
                dtg = r.gt_stamp - prev.gt_stamp
            else:
                dtg = NaN()

            if (
                finite(dtg)
                and dtg > 0
                and finite(r.gt_x)
                and finite(prev.gt_x)
                and finite(r.gt_y)
                and finite(prev.gt_y)
            ):
                dx = r.gt_x - prev.gt_x
                dy = r.gt_y - prev.gt_y
                vx_world = dx / dtg
                vy_world = dy / dtg
                gt_speed = math.hypot(dx, dy) / dtg
                gt_dyaw = wrap_angle(r.gt_yaw - prev.gt_yaw)
                gt_wz = gt_dyaw / dtg

                c = math.cos(r.gt_yaw)
                s = math.sin(r.gt_yaw)
                gt_vx_body = c * vx_world + s * vy_world
                gt_vy_body = -s * vx_world + c * vy_world
            else:
                dtg = gt_speed = gt_wz = gt_vx_body = gt_vy_body = NaN()

            # Absolute errors; useful but not start-aligned.
            if finite(r.gt_x) and finite(r.odom_x):
                ex = r.odom_x - r.gt_x
                ey = r.odom_y - r.gt_y
                ep = math.hypot(ex, ey)
                eyaw = math.degrees(wrap_angle(r.odom_yaw - r.gt_yaw))
            else:
                ex = ey = ep = eyaw = NaN()

            # Timestamp age:
            # Only compare ROS stamps with ROS stamps. Wall time is not mixed
            # into this calculation.
            # The latest ROS stamp age is intentionally not called latency
            # because receipt time and ROS time may use different clocks.
            if finite(r.odom_stamp) and finite(r.joint_stamp):
                stamp_difference_odom_joint_ms = (
                    r.odom_stamp - r.joint_stamp
                ) * 1000.0
            else:
                stamp_difference_odom_joint_ms = NaN()

            if finite(r.gt_stamp) and finite(r.joint_stamp):
                stamp_difference_gt_joint_ms = (
                    r.gt_stamp - r.joint_stamp
                ) * 1000.0
            else:
                stamp_difference_gt_joint_ms = NaN()

            out.append({
                "sample_index": i,
                "wall_monotonic_s": r.wall_time,
                "phase": r.phase,

                "gt_stamp_s": r.gt_stamp,
                "gt_x_m": r.gt_x,
                "gt_y_m": r.gt_y,
                "gt_yaw_deg": math.degrees(r.gt_yaw) if finite(r.gt_yaw) else NaN(),

                "odom_stamp_s": r.odom_stamp,
                "odom_x_m": r.odom_x,
                "odom_y_m": r.odom_y,
                "odom_yaw_deg": math.degrees(r.odom_yaw) if finite(r.odom_yaw) else NaN(),
                "odom_vx_mps": r.odom_vx,
                "odom_wz_rps": r.odom_wz,

                "ekf_stamp_s": r.ekf_stamp,
                "ekf_x_m": r.ekf_x,
                "ekf_y_m": r.ekf_y,
                "ekf_yaw_deg": math.degrees(r.ekf_yaw) if finite(r.ekf_yaw) else NaN(),
                "ekf_vx_mps": r.ekf_vx,
                "ekf_wz_rps": r.ekf_wz,

                "amcl_stamp_s": r.amcl_stamp,
                "amcl_x_m": r.amcl_x,
                "amcl_y_m": r.amcl_y,
                "amcl_yaw_deg": math.degrees(r.amcl_yaw) if finite(r.amcl_yaw) else NaN(),

                "imu_stamp_s": r.imu_stamp,
                "imu_yaw_deg": math.degrees(r.imu_yaw) if finite(r.imu_yaw) else NaN(),
                "imu_wz_rps": r.imu_wz,

                "joint_stamp_s": r.joint_stamp,
                "drive_pos_rad": r.drive_pos,
                "drive_vel_radps": r.drive_vel,
                "steer_pos_rad": r.steer_pos,
                "steer_vel_radps": r.steer_vel,

                "wheel_cmd_radps": r.wheel_cmd_radps,
                "steer_cmd_rad": r.steer_cmd_rad,

                "joint_dt_s": dtj,
                "drive_delta_rad": dtheta,
                "steer_delta_rad": dsteer,
                "drive_distance_m": distance,
                "encoder_velocity_mps": encoder_v,
                "steering_rate_radps": steer_rate,
                "steering_mid_rad": steer_mid,

                "model_yaw_rate_rps": model_wz,
                "gt_yaw_rate_rps": gt_wz,
                "yaw_rate_error_rps": (
                    model_wz - gt_wz
                    if finite(model_wz) and finite(gt_wz)
                    else NaN()
                ),

                "gt_speed_mps": gt_speed,
                "gt_vx_body_mps": gt_vx_body,
                "gt_vy_body_mps": gt_vy_body,

                "odom_error_x_m": ex,
                "odom_error_y_m": ey,
                "odom_position_error_m": ep,
                "odom_heading_error_deg": eyaw,

                "wheel_velocity_tracking_error_radps": wheel_tracking_error,
                "steering_tracking_error_rad": steering_tracking_error,

                "odom_joint_stamp_difference_ms":
                    stamp_difference_odom_joint_ms,
                "gt_joint_stamp_difference_ms":
                    stamp_difference_gt_joint_ms,
            })

        return out

    def write_sheet(self, wb, name, rows):
        ws = wb.create_sheet(name)
        if not rows:
            return

        headers = list(rows[0].keys())
        ws.append(headers)

        for row in rows:
            ws.append([row[h] for h in headers])

        self.format_sheet(ws)

    def format_sheet(self, ws):
        ws.freeze_panes = "A2"

        for c in range(1, ws.max_column + 1):
            max_len = 0
            for rr in range(1, min(ws.max_row, 150) + 1):
                value = ws.cell(rr, c).value
                if value is not None:
                    max_len = max(max_len, len(str(value)))
            ws.column_dimensions[get_column_letter(c)].width = min(
                max(max_len + 2, 10), 32
            )

    def export(self):
        rows = self.derive()

        output_dir = Path(self.args.output_dir).expanduser()
        output_dir.mkdir(parents=True, exist_ok=True)

        filename = time.strftime(
            "bopt_odometry_diagnostic_v2_%Y%m%d_%H%M%S.xlsx"
        )
        path = output_dir / filename

        wb = Workbook()
        wb.remove(wb.active)

        self.write_sheet(wb, "Raw_Data", rows)

        # Phase summary
        groups = {}
        for r in rows:
            if r["phase"] not in ("not_started", "between_tests", "pre_test", "post_test"):
                groups.setdefault(r["phase"], []).append(r)

        summary = []
        for phase, rr in groups.items():
            def V(key):
                return [x[key] for x in rr if finite(x[key])]

            # Distance from GT and odom over the phase.
            gt_d = 0.0
            od_d = 0.0
            for a, b in zip(rr[:-1], rr[1:]):
                if all(finite(a[k]) and finite(b[k]) for k in ("gt_x_m", "gt_y_m")):
                    gt_d += math.hypot(
                        b["gt_x_m"] - a["gt_x_m"],
                        b["gt_y_m"] - a["gt_y_m"],
                    )
                if all(finite(a[k]) and finite(b[k]) for k in ("odom_x_m", "odom_y_m")):
                    od_d += math.hypot(
                        b["odom_x_m"] - a["odom_x_m"],
                        b["odom_y_m"] - a["odom_y_m"],
                    )

            summary.append({
                "phase": phase,
                "samples": len(rr),
                "gt_distance_m": gt_d,
                "odom_distance_m": od_d,
                "distance_error_m": od_d - gt_d,
                "mean_odom_position_error_m": mean(V("odom_position_error_m")),
                "max_odom_position_error_m": max_abs(V("odom_position_error_m")),
                "final_odom_position_error_m": V("odom_position_error_m")[-1]
                    if V("odom_position_error_m") else NaN(),
                "mean_odom_heading_error_deg": mean(V("odom_heading_error_deg")),
                "max_odom_heading_error_deg": max_abs(V("odom_heading_error_deg")),
                "final_odom_heading_error_deg": V("odom_heading_error_deg")[-1]
                    if V("odom_heading_error_deg") else NaN(),
                "mean_yaw_rate_error_rps": mean(V("yaw_rate_error_rps")),
                "max_abs_yaw_rate_error_rps": max_abs(V("yaw_rate_error_rps")),
                "mean_gt_vy_body_mps": mean(V("gt_vy_body_mps")),
                "max_abs_gt_vy_body_mps": max_abs(V("gt_vy_body_mps")),
                "mean_wheel_tracking_error_radps":
                    mean(V("wheel_velocity_tracking_error_radps")),
                "max_abs_wheel_tracking_error_radps":
                    max_abs(V("wheel_velocity_tracking_error_radps")),
                "mean_steering_tracking_error_rad":
                    mean(V("steering_tracking_error_rad")),
                "max_abs_steering_tracking_error_rad":
                    max_abs(V("steering_tracking_error_rad")),
                "mean_joint_dt_s": mean(V("joint_dt_s")),
                "min_joint_dt_s": min(V("joint_dt_s")) if V("joint_dt_s") else NaN(),
                "max_joint_dt_s": max(V("joint_dt_s")) if V("joint_dt_s") else NaN(),
            })

        self.write_sheet(wb, "Phase_Summary", summary)

        # Encoder diagnostics
        enc = []
        for phase, rr in groups.items():
            def V(key):
                return [x[key] for x in rr if finite(x[key])]

            enc.append({
                "phase": phase,
                "samples": len(rr),
                "mean_encoder_velocity_mps": mean(V("encoder_velocity_mps")),
                "rms_encoder_velocity_mps": rms(V("encoder_velocity_mps")),
                "max_abs_encoder_velocity_mps": max_abs(V("encoder_velocity_mps")),
                "mean_drive_delta_rad": mean(V("drive_delta_rad")),
                "max_abs_drive_delta_rad": max_abs(V("drive_delta_rad")),
                "mean_steering_rate_radps": mean(V("steering_rate_radps")),
                "max_abs_steering_rate_radps": max_abs(V("steering_rate_radps")),
                "mean_joint_dt_s": mean(V("joint_dt_s")),
                "min_joint_dt_s": min(V("joint_dt_s")) if V("joint_dt_s") else NaN(),
                "max_joint_dt_s": max(V("joint_dt_s")) if V("joint_dt_s") else NaN(),
                "mean_wheel_tracking_error_radps":
                    mean(V("wheel_velocity_tracking_error_radps")),
                "max_abs_wheel_tracking_error_radps":
                    max_abs(V("wheel_velocity_tracking_error_radps")),
            })

        self.write_sheet(wb, "Encoder_Diagnostics", enc)

        # Curvature diagnostics
        curvature = []
        for phase, rr in groups.items():
            def V(key):
                return [x[key] for x in rr if finite(x[key])]

            curvature.append({
                "phase": phase,
                "samples": len(rr),
                "mean_model_yaw_rate_rps": mean(V("model_yaw_rate_rps")),
                "mean_gt_yaw_rate_rps": mean(V("gt_yaw_rate_rps")),
                "mean_yaw_rate_error_rps": mean(V("yaw_rate_error_rps")),
                "rms_yaw_rate_error_rps": rms(V("yaw_rate_error_rps")),
                "max_abs_yaw_rate_error_rps": max_abs(V("yaw_rate_error_rps")),
                "mean_gt_vy_body_mps": mean(V("gt_vy_body_mps")),
                "rms_gt_vy_body_mps": rms(V("gt_vy_body_mps")),
                "max_abs_gt_vy_body_mps": max_abs(V("gt_vy_body_mps")),
                "mean_steering_tracking_error_rad":
                    mean(V("steering_tracking_error_rad")),
                "max_abs_steering_tracking_error_rad":
                    max_abs(V("steering_tracking_error_rad")),
            })

        self.write_sheet(wb, "Curvature_Diagnostics", curvature)

        # Calibration
        calibration = []

        for phase, rr in groups.items():
            if len(rr) < 2:
                continue

            first = rr[0]
            last = rr[-1]

            if not all(finite(first[k]) and finite(last[k]) for k in
                       ("drive_pos_rad", "gt_x_m", "gt_y_m")):
                continue

            dtheta = last["drive_pos_rad"] - first["drive_pos_rad"]
            gt_distance = math.hypot(
                last["gt_x_m"] - first["gt_x_m"],
                last["gt_y_m"] - first["gt_y_m"],
            )

            reff = (
                gt_distance / abs(dtheta)
                if abs(dtheta) > 1e-8
                else NaN()
            )

            yaw_change = (
                math.radians(last["gt_yaw_deg"] - first["gt_yaw_deg"])
                if finite(last["gt_yaw_deg"]) and finite(first["gt_yaw_deg"])
                else NaN()
            )
            yaw_change = wrap_angle(yaw_change)

            steering = [
                x["steering_mid_rad"]
                for x in rr
                if finite(x["steering_mid_rad"])
            ]
            mean_steering = mean(steering)

            if (
                finite(gt_distance)
                and abs(yaw_change) > 1e-8
                and finite(mean_steering)
                and abs(math.sin(mean_steering)) > 1e-6
            ):
                leff = -gt_distance * math.sin(mean_steering) / yaw_change
            else:
                leff = NaN()

            calibration.append({
                "phase": phase,
                "encoder_delta_rad": dtheta,
                "gt_distance_m": gt_distance,
                "effective_wheel_radius_m": reff,
                "mean_steering_rad": mean_steering,
                "gt_yaw_change_rad": yaw_change,
                "estimated_effective_wheelbase_m": leff,
            })

        self.write_sheet(wb, "Calibration", calibration)

        # Configuration
        ws = wb.create_sheet("Configuration")
        config = [
            ("Version", "V2"),
            ("Wheel radius nominal [m]", self.R),
            ("Wheelbase nominal [m]", self.L),
            ("Drive joint", self.drive_joint),
            ("Steering joint", self.steer_joint),
            ("JointState topic", self.args.joint_topic),
            ("Odom topic", self.args.odom_topic),
            ("EKF topic", self.args.ekf_topic),
            ("AMCL topic", self.args.amcl_topic),
            ("IMU topic", self.args.imu_topic),
            ("Ground truth topic", self.args.gt_topic),
            ("Ground truth type", self.args.gt_type),
            ("Wheel command topic", self.args.wheel_cmd_topic),
            ("Steering command topic", self.args.steer_cmd_topic),
            ("Logger sample rate [Hz]", self.args.sample_hz),
            ("Note", "Command topics are angular wheel velocity [rad/s] and steering angle [rad]."),
            ("Note", "GT-based calibration is valid only if GT topic is correctly configured."),
            ("Note", "Timestamp differences are ROS-stamp differences, not wall-clock latency."),
        ]
        ws.append(["Parameter", "Value"])
        for k, v in config:
            ws.append([k, v])
        self.format_sheet(ws)

        wb.save(path)
        self.get_logger().info(f"Saved: {path}")
        return path


def run_tests(node, args):
    tests = [
        ("stationary", args.stationary_sec,
         "Keep the robot stationary."),
        ("forward_straight", args.motion_sec,
         "Drive FORWARD straight at steady speed."),
        ("reverse_straight", args.motion_sec,
         "Drive REVERSE straight at similar speed magnitude."),
        ("forward_left", args.motion_sec,
         "Drive FORWARD with steady LEFT steering."),
        ("forward_right", args.motion_sec,
         "Drive FORWARD with steady RIGHT steering."),
        ("reverse_left", args.motion_sec,
         "Drive REVERSE with steady LEFT steering."),
        ("reverse_right", args.motion_sec,
         "Drive REVERSE with steady RIGHT steering."),
        ("forward_reverse_transition", args.transition_sec,
         "Perform FORWARD -> STOP -> REVERSE."),
        ("reverse_forward_transition", args.transition_sec,
         "Perform REVERSE -> STOP -> FORWARD."),
    ]

    print("\n============================================================")
    print(" BOPT ODOMETRY DIAGNOSTIC LOGGER V2")
    print("============================================================")
    print("This logger records data only. It does NOT command the robot.")
    print("")

    input("Press ENTER when all ROS nodes are running...")

    node.set_phase("pre_test")
    start = time.monotonic()
    while time.monotonic() - start < args.pre_test_sec:
        rclpy.spin_once(node, timeout_sec=0.02)

    for phase, duration, instruction in tests:
        print("\n------------------------------------------------------------")
        print(f"PHASE: {phase}")
        print(f"{instruction}")
        print(f"Duration: {duration:.1f} s")
        print("------------------------------------------------------------")

        input("Press ENTER to start this phase...")

        node.set_phase(phase)
        start = time.monotonic()

        while time.monotonic() - start < duration and rclpy.ok():
            rclpy.spin_once(node, timeout_sec=0.02)

        node.set_phase("between_tests")
        input("Stop/prepare the robot, then press ENTER for next phase...")

    node.set_phase("post_test")

    start = time.monotonic()
    while time.monotonic() - start < args.post_test_sec:
        rclpy.spin_once(node, timeout_sec=0.02)

    print("\nAll phases completed.")


def main():
    p = argparse.ArgumentParser()

    p.add_argument("--sample-hz", type=float, default=50.0)

    p.add_argument("--wheel-radius", type=float, default=0.115)
    p.add_argument("--wheelbase", type=float, default=1.542)

    p.add_argument("--drive-joint", default="drive_wheel_joint")
    p.add_argument("--steer-joint", default="drive_wheel_Ass_joint")

    p.add_argument("--joint-topic", default="/joint_states")
    p.add_argument("--odom-topic", default="/odom")
    p.add_argument("--ekf-topic", default="/odometry/filtered")
    p.add_argument("--amcl-topic", default="/amcl_pose")
    p.add_argument("--imu-topic", default="/imu/corrected")

    p.add_argument("--gt-topic", default="/ground_truth/pose")
    p.add_argument(
        "--gt-type",
        choices=["pose_stamped", "odometry"],
        default="pose_stamped",
        help="Message type used by the ground-truth topic.",
    )

    p.add_argument(
        "--wheel-cmd-topic",
        default="/traction_joint_controller/commands",
    )
    p.add_argument(
        "--steer-cmd-topic",
        default="/steering_joint_controller/commands",
    )

    p.add_argument("--stationary-sec", type=float, default=8.0)
    p.add_argument("--motion-sec", type=float, default=10.0)
    p.add_argument("--transition-sec", type=float, default=10.0)
    p.add_argument("--pre-test-sec", type=float, default=2.0)
    p.add_argument("--post-test-sec", type=float, default=2.0)

    p.add_argument(
        "--output-dir",
        default="~/bopt_ws/localization_results",
    )

    args = p.parse_args()

    rclpy.init()
    node = DiagnosticNode(args)

    try:
        run_tests(node, args)
        path = node.export()

        print("\n============================================================")
        print(" COMPLETE")
        print("============================================================")
        print(f"Excel: {path}")
        print("")
        print("Review these sheets first:")
        print("  1. Encoder_Diagnostics")
        print("  2. Calibration")
        print("  3. Curvature_Diagnostics")
        print("  4. Phase_Summary")
        print("  5. Raw_Data")
        print("============================================================")

    except KeyboardInterrupt:
        print("\nInterrupted. Saving partial data...")
        try:
            path = node.export()
            print(f"Partial Excel: {path}")
        except Exception as e:
            print(f"Save failed: {e}")
    finally:
        node.destroy_node()
        rclpy.shutdown()


if __name__ == "__main__":
    main()
