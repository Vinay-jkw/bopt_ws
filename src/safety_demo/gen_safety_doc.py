#!/usr/bin/env python3
"""Generate the 360degree Safety Module feature document (.docx) for the testing team."""

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# --- Base style: clean, readable ---
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)

# --- Title ---
title = doc.add_heading("360degree Safety Module", level=0)
title.alignment = WD_ALIGN_PARAGRAPH.LEFT

sub = doc.add_paragraph("Feature Overview for the Testing Team")
sub.runs[0].italic = True
sub.runs[0].font.color.rgb = RGBColor(0x59, 0x59, 0x59)


def heading(text):
    doc.add_heading(text, level=1)


def para(text):
    doc.add_paragraph(text)


def bullets(items):
    for it in items:
        doc.add_paragraph(it, style="List Bullet")


def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = h
        for r in cell.paragraphs[0].runs:
            r.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
    doc.add_paragraph()  # spacing


# --- 1. Purpose ---
heading("1. Purpose")
para(
    "The 360degree Safety Module is a LiDAR-based safety monitoring system for an "
    "autonomous mobile robot (AMR). It continuously watches the area around the robot "
    "using multiple LiDAR sensors, decides whether the surroundings are safe, and "
    "automatically slows or stops the robot when an obstacle enters a protected zone "
    "or when a sensor or connection fault is detected."
)

# --- 2. Safety Zones ---
heading("2. Safety Zones")
para(
    "Each LiDAR sensor defines two protective zones around the robot. The module counts "
    "how many LiDAR points fall inside each zone to decide the safety level. A zone is "
    "considered breached when 5 or more points are detected inside it."
)
bullets([
    "Warning Zone (outer): an obstacle here causes the robot to slow down.",
    "Danger Zone (inner): an obstacle here causes the robot to stop immediately.",
])
para(
    "Zone size adapts automatically to the robot's speed and direction of travel "
    "(forward / reverse and turning direction) — faster motion uses larger zones."
)

# --- 3. Safety Status Levels ---
heading("3. Safety Status Levels")
para(
    "The module reports one overall status at all times. Statuses are evaluated in "
    "priority order (top of the table is highest priority)."
)
table(
    ["Status", "When it happens", "Effect on robot speed"],
    [
        ["Broker Disconnected", "Loss of MQTT connection (no signal for over 2 seconds, or an explicit disconnect message)", "Full stop"],
        ["Danger", "Obstacle detected in a danger zone, or a LiDAR sensor stops responding for over 5 seconds", "Full stop"],
        ["Warning", "Obstacle detected in a warning zone", "Speed reduced to one-third"],
        ["Safe", "No obstacles detected and all systems healthy", "Normal speed (unchanged)"],
    ],
)

# --- 4. Operational Modes ---
heading("4. Operational Modes")
para(
    "The module behaves differently depending on the robot's current operation. The mode "
    "is set externally and changes which sensors are enforced."
)
table(
    ["Mode", "Behavior"],
    [
        ["Normal", "All LiDAR sensors are active and enforced."],
        ["Pickdrop", "Selected side/rear sensors (LiDAR IDs 3, 5, and 6) are temporarily ignored so the robot can approach pick/drop points."],
        ["Parking", "Pickdrop exceptions are bypassed; normal full enforcement applies."],
    ],
)

# --- 5. Safety Turn-Off ---
heading("5. Safety Turn-Off")
para(
    "The safety enforcement can be switched off via an external command. While turned off, "
    "the robot's speed commands pass through unchanged regardless of obstacles. Sending any "
    "other value re-enables enforcement. This is intended for controlled testing and "
    "maintenance situations only."
)

# --- 6. Fault Monitoring (Watchdogs) ---
heading("6. Fault Monitoring")
para(
    "Two independent watchdogs run continuously (checking every 200 milliseconds) to make "
    "sure the robot fails safe if something stops working."
)
bullets([
    "LiDAR watchdog — if any LiDAR sensor goes silent for more than 5 seconds, the status "
    "becomes Danger (robot stops) and error code E011 is reported.",
    "MQTT watchdog — if the MQTT connection is lost for more than 2 seconds, the status "
    "becomes Broker Disconnected (robot stops).",
])

heading("7. Error Codes")
table(
    ["Code", "Meaning"],
    [
        ["E011", "LiDAR data timeout — no data received from a sensor for over 5 seconds."],
    ],
)

# --- 8. Configuration Input (Database) ---
heading("8. Configuration Input (Safety Policy Database)")
para(
    "The module does not hard-code its safety behavior. All LiDAR settings and safety zone "
    "definitions are loaded at startup from a single SQLite database file (the default file "
    "is config/bopt_2000.db). To test a different configuration, swap in a different database "
    "file. The database holds two kinds of records: one LiDAR configuration per sensor, and "
    "a set of policies that define the zones for each sensor."
)

doc.add_heading("8.1 LiDAR Configuration", level=2)
para(
    "Describes each LiDAR sensor — where it is mounted on the robot and how its zone is "
    "shaped and oriented."
)
table(
    ["Field", "Meaning"],
    [
        ["lidar_id", "Numeric ID for the sensor, used throughout the system."],
        ["topic", "ROS topic the sensor publishes its scan on."],
        ["shape_type", "Zone shape: Rectangle, L-Shape, or Mirror L-Shape."],
        ["x_offset, y_offset", "Sensor position relative to the robot base."],
        ["theta, theta_N", "Rotation angles used to orient the zone."],
        ["description", "Human-readable label for the sensor."],
    ],
)

doc.add_heading("8.2 Safety Policy", level=2)
para(
    "Each policy defines the warning and danger zones that apply to one LiDAR under a "
    "specific speed and direction of travel. A LiDAR has many policies; the module "
    "automatically picks the matching policy for the robot's current motion (it selects "
    "the tightest policy whose maximum speed is at or above the current speed, with a "
    "matching direction)."
)
table(
    ["Field", "Meaning"],
    [
        ["id", "Policy ID."],
        ["field_type", "Which zone this policy defines: Warning or Danger."],
        ["max_speed", "Linear speed up to which this policy applies."],
        ["angular_max_speed", "Angular (turning) speed up to which this policy applies."],
        ["direction", "Travel direction this policy is for: Forward or Reverse."],
        ["angular_direction", "Turning direction: Clockwise or Counter-Clockwise."],
        ["warning_zone / danger_zone", "Zone dimensions, made up of the parameters below."],
        ["description", "Human-readable label for the policy."],
    ],
)
para("Each zone (warning or danger) is described by these dimension parameters:")
table(
    ["Parameter", "Meaning"],
    [
        ["a", "Width parameter."],
        ["b", "Length parameter."],
        ["i1, i2", "Inner shape parameters."],
        ["o1, o2", "Outer shape parameters."],
    ],
)

# --- 9. Visualization ---
heading("9. Visualization")
para(
    "The module publishes the safety zones as overlays that can be viewed live in RViz, so "
    "testers can visually confirm zone shape and size. Warning zones are drawn in yellow and "
    "danger zones in red, one set per LiDAR sensor. A ready-made RViz layout (safety_viz.rviz) "
    "is included."
)
para(
    "Markers are published per LiDAR sensor, in that sensor's own frame (frame_<topic>). "
    "For a sensor with the given LiDAR ID:"
)
table(
    ["Marker topic", "Zone", "Appearance"],
    [
        ["visualization_marker_<id>_warning", "Warning zone", "Yellow line outline (line width 0.03 m)"],
        ["visualization_marker_<id>_danger", "Danger zone", "Red line outline (line width 0.05 m)"],
    ],
)

# --- 10. Interfaces for Testing ---
heading("10. Interfaces for Testing")
para("Inputs the module listens to (use these to drive test scenarios):")
table(
    ["Signal", "Purpose"],
    [
        ["LiDAR scans", "Obstacle data from each sensor (topics are loaded from configuration)."],
        ["/odometry/filtered", "Robot speed and direction."],
        ["safety_turnoff", "Send \"TurnOff\" to disable enforcement; any other value re-enables it."],
        ["/byd/safety", "Set operational mode: \"normal\", \"pickdrop\", or \"parking\"."],
        ["/mqtt/status", "Connection health: \"connected\" or \"disconnected\"."],
    ],
)
para("Outputs the module produces (observe these to verify behavior):")
table(
    ["Signal", "Purpose"],
    [
        ["safety_status", "Current status: safe, warning, danger, or broker_disconnected."],
        ["/cmd_vel_remapped", "Robot velocity command after safety scaling."],
        ["/velocity_remapped", "Scalar velocity after safety scaling."],
        ["machine/error/status", "Active error codes (e.g. E011)."],
        ["visualization_marker_<id>_warning / _danger", "Warning (yellow) and danger (red) zone overlays for RViz, per LiDAR."],
    ],
)

out = "/media/aparnesh/D/CodeRefactoring/bopt_v2_ws/src/safety_demo/360degree_Safety_Module.docx"
doc.save(out)
print("saved:", out)
